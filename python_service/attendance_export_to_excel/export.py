from aiohttp import web
import pandas as pd
import tempfile
from openpyxl import load_workbook
from openpyxl.styles import Font
from s3_client.aws_utils import upload_and_presign

REPLACEMENTS = {
    "work_daytime": "Я",                    # Дневная работа
    "work_nighttime": "Н",                  # Ночная работа
    "work_weekend_holiday": "РВ",           # Работа в выходные/праздники
    "overtime": "C",                        # Сверхурочные
    "shortened_work_time": "ЛЧ",            # Сокращенное рабочее время
    "part_time": "ЛЧ",
    "vacation_days_instead": "ОТ",

    "vacation": "ОТ",                       # Отпуск
    "additional_paid_vacation": "ОД",       # Дополнительный оплачиваемый отпуск
    "unpaid_vacation": "ДО",                # Отпуск без содержания
    "unpaid_vacation_with_reason": "ДО",

    "study_vacation_paid": "У",             # Учебный отпуск (оплачиваемый)
    "study_vacation_unpaid": "УД",          # Учебный отпуск (неоплачиваемый)

    "maternity_leave": "Р",                 # Декретный отпуск
    "childcare_leave": "ОЖ",                # Отпуск по уходу за ребенком

    "sick_leave": "Б",                      # Больничный

    "business_trip": "K",                   # Командировка

    "absence_without_reason": "ПР",         # Прогул
    "absence_unknown_reason": "НН",         # Отсутствие по неизвестной причине

    "part_time_by_employer": "НС",          # Неполное время по инициативе работодателя
    "weekend_holiday": "В",                 # Выходной/праздник
    "downtime_not_employer_fault": "НП",    # Простой не по вине работодателя
    "suspension_with_pay": "НО",            # Отстранение с оплатой

    "payout_birth": "",
    "certificates_on_dismissal": "",
    "tax_deduction_children": "",
    "maternity_childcare_15": "ОЖ",
    "maternity_childcare_3": "ОЖ",
    "maternity_pregnancy": "Р",
    "transfer": "",
    "vacation_shift": "ОТ",
    "maternity_early_exit": "",
    "resignation": "",
    "maternity_work_during": "Р",
    "personal_data_change": "",

    "tracker_task_deadline": "ТД",
}

WORKING_DAYS = (
    "work_daytime",
    "work_nighttime",
    "work_weekend_holiday",
    "overtime",
    "shortened_work_time",
    "part_time_by_employer"
)

def transform_attendance(attendances: list[dict], date_from: str, date_to: str) -> dict[str, pd.DataFrame]:
    df = pd.DataFrame(attendances)

    # ФИО + компания
    df["ФИО"] = df["employee"].apply(
        lambda e: f"{e.get('surname','')} {e.get('name','')} {e.get('patronymic','')}".strip()
    )
    df["Компания"] = df["employee"].apply(lambda e: e.get("subcompany", ""))

    # Даты (могут быть пустыми)
    df["start_date"] = pd.to_datetime(df.get("start_date"))
    df["end_date"] = pd.to_datetime(df.get("end_date"))

    # приоритет: abscence_type > attendance_type

    def pick_code(r):
        att = r.get("attendance_type")
        absn = r.get("abscence_type")

        att  = att  if pd.notna(att)  and att  != "" else None
        absn = absn if pd.notna(absn) and absn != "" else None
        return (absn or att or "")


    df["raw_value"] = df.apply(pick_code, axis=1)

    # замена по словарю
    df["value"] = df["raw_value"].map(lambda v: REPLACEMENTS.get(v, v))

    # часы только для рабочих смен
    df["hours"] = df.apply(
        lambda r: ((r["end_date"] - r["start_date"]).total_seconds() / 3600.0)
        if pd.notna(r["start_date"]) and pd.notna(r["end_date"]) and r.get("attendance_type") in WORKING_DAYS
        else 0.0,
        axis=1
    )

    all_days = pd.date_range(date_from, date_to, freq="D")

    company_tables: dict[str, pd.DataFrame] = {}
    for company, group in df.groupby("Компания"):
        rows = []
        for _, r in group.iterrows():
            if pd.isna(r["start_date"]) and pd.isna(r["end_date"]):
                for day in all_days:
                    rows.append({
                        "ФИО": r["ФИО"],
                        "day": day,
                        "value": "",
                        "label": "",
                        "hours": 0.0,
                    })
                continue

            start_norm = (
                r["start_date"].normalize()
                if pd.notna(r["start_date"]) else r["end_date"].normalize()
            )
            end_norm = (
                r["end_date"].normalize()
                if pd.notna(r["end_date"]) else start_norm
            )
            for day in pd.date_range(start_norm, end_norm, freq="D"):
                hours_day = r["hours"] if day == start_norm else 0.0
                label = f"{r['value']} {hours_day:.2f}".rstrip() if hours_day > 0 else (r["value"] or "")

                rows.append({
                    "ФИО": r["ФИО"],
                    "day": day,
                    "value": r["value"],
                    "label": label,
                    "hours": hours_day,
                })

        df_days = pd.DataFrame(rows)
        if df_days.empty:
            continue

        # сводная таблица: по дням — значения
        pivot = df_days.pivot_table(
            index="ФИО",
            columns="day",
            values="label",
            aggfunc=lambda x: ", ".join(str(v) for v in x if pd.notna(v) and v != "")
        ).fillna("")

        # добавляем все дни периода
        pivot = pivot.reindex(columns=all_days, fill_value="")

        # меняем названия колонок на числа
        day_labels = {d: str(d.day) for d in all_days}
        pivot.rename(columns=day_labels, inplace=True)

        pivot.reset_index(inplace=True)

        # считаем часы
        work_hours = df_days.groupby("ФИО")["hours"].sum()
        pivot["work_hours"] = pivot["ФИО"].map(work_hours).fillna(0.0).round(2)

        ordered_cols = ["ФИО"] + [str(d.day) for d in all_days] + ["work_hours"]
        pivot = pivot.reindex(columns=ordered_cols, fill_value="")

        company_tables[company] = pivot

    return company_tables


async def generate_attendance_excel(request):
    file_key = request.rel_url.query['file_key']
    from_date = request.rel_url.query['from_date']
    to_date = request.rel_url.query['to_date']

    if not from_date or not to_date:
        return web.Response(status=400, text="from_date and to_date are required")

    data = await request.json()
    attendances = data.get("attendances", [])

    if not attendances:
        return web.Response(status=400, text="Invalid input data")

    company_tables = transform_attendance(attendances, from_date, to_date)

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        output = tmp.name

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        row_offset = 0
        period_text = f"Период: {from_date} — {to_date}"
        pd.DataFrame([[period_text]]).to_excel(
            writer, sheet_name="Attendance", startrow=row_offset, index=False, header=False
        )
        row_offset += 2
        for company, table in company_tables.items():
            pd.DataFrame([[f"Компания {company}"]]).to_excel(
                writer, sheet_name="Attendance", startrow=row_offset, index=False, header=False
            )
            row_offset += 1

            table.to_excel(
                writer, sheet_name="Attendance", startrow=row_offset, index=False
            )
            row_offset += len(table) + 2

    wb = load_workbook(output)
    ws = wb["Attendance"]
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        cell = row[0]
        if cell.value and str(cell.value).startswith("Компания "):
            cell.font = Font(bold=True, size=12)
    wb.save(output)

    url = upload_and_presign(output, file_key + ".xlsx")
    return web.Response(status=200, content_type="text/plain", text=url)
